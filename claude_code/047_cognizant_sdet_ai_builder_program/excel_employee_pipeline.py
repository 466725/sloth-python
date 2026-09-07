"""Excel -> dicts -> filter -> JSON -> summary-sheet pipeline (openpyxl practice)."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

# Demo artifacts go under temps/ per repo convention (generated outputs).
OUTPUT_DIR = Path("temps/excel_employee_pipeline")
EXCEL_PATH = OUTPUT_DIR / "employees.xlsx"
JSON_PATH = OUTPUT_DIR / "high_earners.json"

HEADERS = ["EmpID", "Name", "Department", "Salary"]
SAMPLE_ROWS = [
    ["E001", "Alice Chen", "Engineering", 92000],
    ["E002", "Bob Patel", "Engineering", 48000],
    ["E003", "Carol Diaz", "Sales", 61000],
    ["E004", "Dan Kim", "Sales", 45000],
    ["E005", "Eve Novak", "HR", 53000],
]


def create_sample_excel(path: Path) -> None:
    """(1) Write a sample workbook with the employee table."""
    wb = Workbook()
    try:
        ws = wb.active
        ws.title = "Employees"
        ws.append(HEADERS)
        for row in SAMPLE_ROWS:
            ws.append(row)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    finally:
        # finally guarantees the workbook is closed even if save() fails.
        wb.close()


def read_employees(path: Path) -> list[dict]:
    """(2) Read rows back as a list of dicts keyed by the header row."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        # Re-raise with context: the caller sees WHICH file was missing.
        raise FileNotFoundError(f"Employee workbook not found: {path}") from None

    try:
        ws = wb["Employees"]
        # values_only=True yields plain tuples; zip() pairs each row with the
        # headers to build dicts without index juggling.
        rows = list(ws.iter_rows(values_only=True))
        header, *data_rows = rows
        return [dict(zip(header, row)) for row in data_rows]
    finally:
        wb.close()


def filter_high_earners(employees: list[dict], threshold: float) -> list[dict]:
    """(3) Keep employees whose Salary exceeds the threshold."""
    return [emp for emp in employees if emp["Salary"] > threshold]


def write_json(employees: list[dict], path: Path) -> None:
    """(4) Write filtered rows to JSON with stable, human-readable formatting."""
    path.parent.mkdir(parents=True, exist_ok=True)
    # indent=2 + sorted keys: diff-friendly output for reviews and test asserts.
    path.write_text(
        json.dumps(employees, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def summarize_by_department(employees: list[dict]) -> dict[str, dict[str, float]]:
    """Count employees and average salary per department."""
    summary: dict[str, dict[str, float]] = {}
    for emp in employees:
        dept = summary.setdefault(emp["Department"], {"count": 0, "total": 0.0})
        dept["count"] += 1
        dept["total"] += emp["Salary"]
    for dept in summary.values():
        dept["avg_salary"] = dept.pop("total") / dept["count"]
    return summary


def write_summary_sheet(path: Path, summary: dict[str, dict[str, float]]) -> None:
    """(5) Append a 'Summary' sheet with count + avg salary per department."""
    try:
        wb = load_workbook(path)
    except FileNotFoundError:
        raise FileNotFoundError(f"Cannot add summary, workbook missing: {path}") from None

    try:
        if "Summary" in wb.sheetnames:
            del wb["Summary"]  # idempotent re-runs: replace, don't duplicate
        ws = wb.create_sheet("Summary")
        ws.append(["Department", "Count", "AvgSalary"])
        for department, stats in sorted(summary.items()):
            ws.append([department, stats["count"], round(stats["avg_salary"], 2)])
        wb.save(path)
    finally:
        wb.close()


if __name__ == "__main__":
    create_sample_excel(EXCEL_PATH)
    print(f"created {EXCEL_PATH}")

    employees = read_employees(EXCEL_PATH)
    print(f"read {len(employees)} employees")

    high_earners = filter_high_earners(employees, threshold=50000)
    write_json(high_earners, JSON_PATH)
    print(f"wrote {len(high_earners)} high earners to {JSON_PATH}")

    summary = summarize_by_department(employees)
    write_summary_sheet(EXCEL_PATH, summary)
    for dept, stats in sorted(summary.items()):
        print(f"  {dept}: count={stats['count']} avg={stats['avg_salary']:.2f}")

    # Error-handling demo: reading a missing file raises our contextual error.
    try:
        read_employees(OUTPUT_DIR / "does_not_exist.xlsx")
    except FileNotFoundError as exc:
        print(f"caught: {exc}")
